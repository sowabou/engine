"""Binance P2P USDT→MRU parser.

Parses Cadorim's Binance P2P sales history (Binance_P2P_USDT_MRU_2026.xlsx)
into a list of canonical exchange events.

Each row = one USDT sale converted to MRU. Used as the "MRUUSDExchange
account" — the implicit FX rate per day comes from these sales.

The output supports two views:
1. Per-transaction list (every sale with date, USDT, MRU, rate)
2. Daily aggregate (total USDT sold, MRU received, weighted avg rate per day)
"""
from __future__ import annotations

import os
from dataclasses import asdict, dataclass
from datetime import datetime
from typing import Iterable

try:
    import openpyxl
except ImportError:
    openpyxl = None


@dataclass
class BinanceSale:
    """One USDT→MRU sale on Binance P2P."""

    date: str
    time: str
    usdt_sold: float
    mru_received: float
    rate: float           # MRU per USDT for this sale
    fee_taker: float
    counterparty: str
    order_id: str

    def to_dict(self) -> dict:
        return asdict(self)


def parse_p2p_file(path: str) -> list[BinanceSale]:
    """Parse Binance_P2P_USDT_MRU_*.xlsx → list of sales."""
    if openpyxl is None or not os.path.exists(path):
        return []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Transactions" not in wb.sheetnames:
        wb.close()
        return []
    ws = wb["Transactions"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    sales: list[BinanceSale] = []
    for r in rows[1:]:
        if not r or r[0] is None:
            continue
        date_val, time_val, usdt, mru, rate, fee, cp, order = (list(r) + [None] * 8)[:8]
        # Parse date
        if isinstance(date_val, datetime):
            date_iso = date_val.strftime("%Y-%m-%d")
        else:
            try:
                date_iso = datetime.strptime(str(date_val), "%d/%m/%Y").strftime("%Y-%m-%d")
            except (ValueError, TypeError):
                continue
        try:
            sale = BinanceSale(
                date=date_iso,
                time=str(time_val) if time_val else "",
                usdt_sold=float(usdt or 0),
                mru_received=float(mru or 0),
                rate=float(rate or 0),
                fee_taker=float(fee or 0),
                counterparty=str(cp or ""),
                order_id=str(order or ""),
            )
        except (ValueError, TypeError):
            continue
        sales.append(sale)
    sales.sort(key=lambda s: (s.date, s.time))
    return sales


def aggregate_by_day(sales: Iterable[BinanceSale]) -> list[dict]:
    """Group sales by date, compute weighted-average rate per day."""
    by_date: dict[str, dict] = {}
    for s in sales:
        d = by_date.setdefault(s.date, {
            "date": s.date, "n_sales": 0, "usdt_sold": 0.0,
            "mru_received": 0.0, "fees": 0.0,
        })
        d["n_sales"] += 1
        d["usdt_sold"] += s.usdt_sold
        d["mru_received"] += s.mru_received
        d["fees"] += s.fee_taker
    daily = []
    for d in by_date.values():
        d["rate_avg"] = (d["mru_received"] / d["usdt_sold"]) if d["usdt_sold"] > 0 else 0.0
        d["rate_avg"] = round(d["rate_avg"], 4)
        d["usdt_sold"] = round(d["usdt_sold"], 2)
        d["mru_received"] = round(d["mru_received"], 2)
        d["fees"] = round(d["fees"], 4)
        daily.append(d)
    daily.sort(key=lambda x: x["date"])
    return daily


def filter_period(sales: list[BinanceSale], period_from: str | None, period_to: str | None) -> list[BinanceSale]:
    out = sales
    if period_from:
        out = [s for s in out if s.date >= period_from]
    if period_to:
        out = [s for s in out if s.date <= period_to]
    return out


if __name__ == "__main__":
    import sys
    workspace = sys.argv[1] if len(sys.argv) > 1 else "/sessions/adoring-keen-clarke/mnt/Cadorim Accounting"
    sales = parse_p2p_file(os.path.join(workspace, "Binance_P2P_USDT_MRU_2026.xlsx"))
    print(f"Parsed {len(sales)} P2P sales")
    if sales:
        print(f"Date range: {sales[0].date} → {sales[-1].date}")
        daily = aggregate_by_day(sales)
        print(f"Days with sales: {len(daily)}")
        print(f"\nFirst 5 daily aggregates:")
        for d in daily[:5]:
            print(f"  {d['date']}: {d['n_sales']:>3} sales, {d['usdt_sold']:>10,.2f} USDT, {d['mru_received']:>12,.2f} MRU, rate {d['rate_avg']}")
